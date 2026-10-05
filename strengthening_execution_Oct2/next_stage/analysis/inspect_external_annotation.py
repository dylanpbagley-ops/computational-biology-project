"""Summarize the publicly retrieved GSE305971 per-cell annotation workbook."""

from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
NEXT = HERE.parent
SOURCE = NEXT / "data" / "GSE305971_cells_annotation_for_geo_ver_20250306.xlsx"
OUT = HERE / "results" / "external_annotation_by_sample.tsv"


def main() -> None:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    column = {name: index for index, name in enumerate(header)}
    required = ("Amp.Batch.ID", "Hospital.Code", "Biopsy.Sequence", "Populations", "Disease", "Tissue", "Method", "Seq.Batch.ID")
    if any(name not in column for name in required):
        raise ValueError(f"Annotation columns missing: {required}")

    summaries: dict[str, dict[str, object]] = defaultdict(lambda: {"cells": 0, "fields": {name: Counter() for name in required[1:]}})
    disease_population_cells: Counter[tuple[str, str, str]] = Counter()
    donors_by_disease: dict[str, set[tuple[str, str]]] = defaultdict(set)
    total = 0
    for row_number, row in enumerate(rows, start=2):
        if len(row) != len(header) or row[column["Biopsy.Sequence"]] is None:
            raise ValueError(f"Malformed annotation row {row_number}")
        sample = str(row[column["Amp.Batch.ID"]])
        entry = summaries[sample]
        entry["cells"] += 1
        total += 1
        values = {}
        for name in required[1:]:
            value = row[column[name]]
            values[name] = "NR" if value is None or str(value).strip() == "" else str(value).strip()
            entry["fields"][name][values[name]] += 1
        disease_population_cells[(values["Disease"], values["Tissue"], values["Populations"])] += 1
        donors_by_disease[values["Disease"]].add((values["Hospital.Code"], values["Biopsy.Sequence"]))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=("Amp.Batch.ID", "cell_count", *(f"{name}_counts" for name in required[1:])),
            delimiter="\t",
            quoting=csv.QUOTE_MINIMAL,
        )
        writer.writeheader()
        for sample in sorted(summaries):
            entry = summaries[sample]
            output = {"Amp.Batch.ID": sample, "cell_count": entry["cells"]}
            for name, counter in entry["fields"].items():
                output[f"{name}_counts"] = json.dumps(dict(sorted(counter.items())), ensure_ascii=False, sort_keys=True)
            writer.writerow(output)

    print(json.dumps({
        "source": str(SOURCE),
        "workbook_sha256": __import__("hashlib").sha256(SOURCE.read_bytes()).hexdigest(),
        "annotation_cells": total,
        "amp_batch_samples": len(summaries),
        "distinct_hospital_biopsy_pairs_by_disease": {disease: len(donors) for disease, donors in sorted(donors_by_disease.items())},
        "cell_counts_by_disease_tissue_population": {"|".join(key): count for key, count in sorted(disease_population_cells.items())},
        "output": str(OUT),
    }, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
