"""Execute the frozen, code-held-out GSE305971 MARS feasibility pilot."""

from __future__ import annotations

import csv
import gzip
import hashlib
import json
import math
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
NEXT = HERE.parent
DATA = NEXT / "data"
RESULTS = HERE / "results"
ANNOTATION = DATA / "GSE305971_cells_annotation_for_geo_ver_20250306.xlsx"
PILOT_SUMMARY = RESULTS / "external_pilot_manifest_locked.json"
PILOT_MANIFEST = RESULTS / "external_pilot_manifest_locked.tsv"
DOWNLOAD_MANIFEST = RESULTS / "external_pilot_download_manifest.tsv"
PROTOCOL = HERE / "selected_protocol_v2.md"
PROTOCOL_HASH = HERE / "selected_protocol_v2.md.sha256"

EXPECTED_WORKBOOK_SHA = "2fc3f9e19ab82097f6a50689de7bf4083aeb81aba442f263ea81afe8b150b94d"
EXPECTED_PILOT_MANIFEST_SHA = "43f17abd39a15136aeaaecb183887752c6a76d38d96e1626014abd60953508cb"
EXPECTED_DOWNLOAD_MANIFEST_SHA = "04f24edf9473468838bd377477ea08590399c9cf1820272d5eaddca44fea5864"
EXPECTED_PROTOCOL_SHA = "aa7b16216786b308bd8fb49f375fb481793fc3169bd3e4719bc9966892a7029b"
PREEXPOSED_CODES = {"tlv53"}
MIN_CELLS = 50
MARKERS = ("TNFRSF17", "GPRC5D", "SLAMF7")
GROUP_MAP = {
    ("Healthy", "Normal_PC"): "healthy_normal_pc",
    ("Healthy", "Malignant"): "healthy_ambiguous_malignant_label",
    ("NDMM", "Malignant"): "mm_malignant",
    ("RRMM", "Malignant"): "mm_malignant",
    ("NDMM", "Normal_PC"): "mm_normal_pc_prior_art_reference",
    ("RRMM", "Normal_PC"): "mm_normal_pc_prior_art_reference",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def write_tsv(path: Path, fields: tuple[str, ...], rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        writer.writerows(rows)


def as_float(value: float | np.floating) -> float:
    value = float(value)
    if not math.isfinite(value):
        raise ValueError(f"Non-finite result: {value}")
    return value


def get_group(disease: str, population: str) -> str | None:
    return GROUP_MAP.get((disease, population))


def read_selected_annotations(selected_codes: set[str], batch_ids: set[str]):
    workbook = load_workbook(ANNOTATION, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise ValueError("Annotation workbook is empty")
    columns = {str(name): i for i, name in enumerate(header)}
    required = (
        "index", "Amp.Batch.ID", "Hospital.Code", "Biopsy.Sequence", "Disease",
        "Tissue", "Method", "Populations", "Seq.Batch.ID",
    )
    missing = sorted(set(required) - set(columns))
    if missing:
        raise ValueError(f"Missing annotation fields: {missing}")
    selected: dict[str, dict[str, str]] = {}
    selected_counts: Counter[tuple[str, str]] = Counter()
    for row_number, row in enumerate(rows, start=2):
        batch_id = str(row[columns["Amp.Batch.ID"]]).strip()
        if batch_id not in batch_ids:
            continue
        cell_id = str(row[columns["index"]]).strip()
        code = str(row[columns["Hospital.Code"]]).strip()
        sequence = str(row[columns["Biopsy.Sequence"]]).strip()
        disease = str(row[columns["Disease"]]).strip()
        tissue = str(row[columns["Tissue"]]).strip()
        method = str(row[columns["Method"]]).strip()
        population = str(row[columns["Populations"]]).strip()
        group = get_group(disease, population)
        if code not in selected_codes or sequence != "1" or tissue != "BM" or method != "MARS" or group is None:
            continue
        if cell_id in selected:
            raise ValueError(f"Duplicate selected annotation cell ID at row {row_number}: {cell_id}")
        selected[cell_id] = {
            "Amp.Batch.ID": batch_id,
            "Hospital.Code": code,
            "Biopsy.Sequence": sequence,
            "Disease": disease,
            "Tissue": tissue,
            "Method": method,
            "Populations": population,
            "Seq.Batch.ID": str(row[columns["Seq.Batch.ID"]]).strip(),
            "analysis_group": group,
        }
        selected_counts[(group, code)] += 1
    workbook.close()
    return selected, selected_counts


def parse_count_file(path: Path, wanted_cell_ids: set[str]):
    """Read a genes-by-cells gzip table with NumPy's C parser and column sums."""
    with gzip.open(path, "rb") as handle:
        header_line = handle.readline().rstrip(b"\r\n")
        cell_ids = [value.decode("utf-8") for value in header_line.split(b"\t")]
        if not cell_ids or len(cell_ids) != len(set(cell_ids)):
            raise ValueError(f"Missing or duplicate cell IDs in {path.name}")
        positions = [i for i, cell_id in enumerate(cell_ids) if cell_id in wanted_cell_ids]
        totals = np.zeros(len(cell_ids), dtype=np.int64)
        marker_values: dict[str, np.ndarray] = {}
        gene_count = 0
        for line_no, line in enumerate(handle, start=2):
            line = line.rstrip(b"\r\n")
            gene, separator, payload = line.partition(b"\t")
            if not separator:
                raise ValueError(f"Malformed row {line_no} in {path.name}")
            gene_name = gene.decode("utf-8")
            values = np.fromstring(payload, dtype=np.int64, sep="\t")
            if values.size != len(cell_ids):
                raise ValueError(
                    f"Malformed count row {line_no} in {path.name}: expected {len(cell_ids)}, got {values.size}"
                )
            if np.any(values < 0):
                raise ValueError(f"Negative count in row {line_no} of {path.name}")
            if np.any(values > np.iinfo(np.int64).max - totals):
                raise ValueError(f"Integer count sum overflow in row {line_no} of {path.name}")
            totals += values
            gene_count += 1
            if gene_name in MARKERS:
                if gene_name in marker_values:
                    raise ValueError(f"Duplicate marker row {gene_name} in {path.name}")
                marker_values[gene_name] = values.copy()
    missing_markers = sorted(set(MARKERS) - set(marker_values))
    if missing_markers:
        raise ValueError(f"Marker rows missing from {path.name}: {missing_markers}")
    return cell_ids, positions, totals, marker_values, gene_count


def make_folds(codes: list[str], n_folds: int, rng: np.random.Generator) -> dict[str, int]:
    if len(codes) < n_folds:
        raise ValueError(f"Need at least {n_folds} codes for {n_folds}-fold evaluation")
    shuffled = list(rng.permutation(np.asarray(sorted(codes), dtype=object)))
    return {str(code): i % n_folds for i, code in enumerate(shuffled)}


def candidate_thresholds(training_by_code: dict[str, np.ndarray]) -> np.ndarray:
    """Shared grid: zero and each observed value, both sides of each tie."""
    if not training_by_code or any(values.size == 0 for values in training_by_code.values()):
        raise ValueError("Threshold training codes must each contain reference values")
    values = np.concatenate(list(training_by_code.values())).astype(np.float64, copy=False)
    if not np.all(np.isfinite(values)) or np.any(values < 0):
        raise ValueError("Reference expression values must be finite and nonnegative")
    observed = np.unique(values)
    above_each = np.nextafter(observed, np.inf)
    return np.unique(np.concatenate((np.asarray([0.0]), observed, above_each)))


def fit_thresholds(
    policy: str,
    training_by_code: dict[str, np.ndarray],
    budget: float,
) -> tuple[float, float]:
    if policy not in {"pooled", "equal_code", "donor_envelope"}:
        raise ValueError(f"Unknown policy: {policy}")
    candidates = candidate_thresholds(training_by_code)
    for candidate in candidates:
        donor_rates = [
            np.count_nonzero(values >= candidate) / values.size for values in training_by_code.values()
        ]
        if policy == "pooled":
            positives = sum(np.count_nonzero(values >= candidate) for values in training_by_code.values())
            total = sum(values.size for values in training_by_code.values())
            rate = positives / total
        elif policy == "equal_code":
            rate = float(np.mean(donor_rates))
        else:
            rate = max(donor_rates)
        if rate <= budget:
            return float(candidate), float(rate)
    raise AssertionError("Shared candidate grid must include a cutoff above every training value")


def classify_threshold(value: float, max_training_value: float) -> str:
    return "ABOVE_MAX" if not math.isfinite(value) or value > max_training_value else "FINITE"


def format_threshold(value: float) -> str:
    return format(float(value), ".17g") if math.isfinite(value) else "INF"


def threshold_positive(values: np.ndarray, threshold: float) -> int:
    return int(np.count_nonzero(values >= threshold))


def exact_batch_matches(cell_ids: list[str], annotations: dict[str, dict[str, str]], batch_id: str):
    """Join count columns to annotations by cell ID, retaining only the stated batch."""
    return [
        (position, cell_id, annotations[cell_id])
        for position, cell_id in enumerate(cell_ids)
        if cell_id in annotations and annotations[cell_id]["Amp.Batch.ID"] == batch_id
    ]


def metric_rows(cells: list[dict[str, object]], selected_codes: set[str]):
    by_group_code: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in cells:
        by_group_code[(str(row["analysis_group"]), str(row["Hospital.Code"]))].append(row)
    code_rows = []
    for group, codes in (
        ("healthy_normal_pc", sorted(selected_codes & {c for g, c in by_group_code if g == "healthy_normal_pc"})),
        ("healthy_ambiguous_malignant_label", sorted(selected_codes & {c for g, c in by_group_code if g == "healthy_ambiguous_malignant_label"})),
        ("mm_malignant", sorted(selected_codes & {c for g, c in by_group_code if g == "mm_malignant"})),
        ("mm_normal_pc_prior_art_reference", sorted(selected_codes & {c for g, c in by_group_code if g == "mm_normal_pc_prior_art_reference"})),
    ):
        for code in codes:
            group_cells = by_group_code[(group, code)]
            for marker in MARKERS:
                raw = np.asarray([int(row[f"{marker}_molecules"]) for row in group_cells], dtype=np.int64)
                norm = np.asarray([float(row[f"{marker}_per_10k"]) for row in group_cells], dtype=np.float64)
                counts = np.asarray([int(row["total_mRNA_molecule_counts"]) for row in group_cells], dtype=np.int64)
                code_rows.append({
                    "analysis_group": group,
                    "Hospital.Code": code,
                    "marker": marker,
                    "n_exact_joined_cells": int(raw.size),
                    "meets_50_cell_sensitivity": raw.size >= MIN_CELLS,
                    "preexposed_code": code in PREEXPOSED_CODES,
                    "median_total_mRNA_molecules": float(np.median(counts)),
                    "mean_raw_molecules": float(np.mean(raw)),
                    "raw_detection_fraction": float(np.count_nonzero(raw) / raw.size),
                    "mean_per_10k": float(np.mean(norm)),
                    "median_per_10k": float(np.median(norm)),
                    "q90_per_10k": float(np.quantile(norm, 0.90)),
                })
    return code_rows, by_group_code


def paired_contrasts(by_group_code: dict[tuple[str, str], list[dict[str, object]]], mm_codes: set[str]):
    rows = []
    for code in sorted(mm_codes):
        malignant = by_group_code.get(("mm_malignant", code), [])
        normal = by_group_code.get(("mm_normal_pc_prior_art_reference", code), [])
        for marker in MARKERS:
            malign = np.asarray([float(r[f"{marker}_per_10k"]) for r in malignant], dtype=np.float64)
            norm = np.asarray([float(r[f"{marker}_per_10k"]) for r in normal], dtype=np.float64)
            if malign.size == 0 or norm.size == 0:
                rows.append({
                    "Hospital.Code": code,
                    "marker": marker,
                    "n_malignant_cells": int(malign.size),
                    "n_within_code_Normal_PC_cells": int(norm.size),
                    "paired_mean_difference_malignant_minus_Normal_PC": "NR",
                    "normal_reference_sd": "NR",
                    "standardized_mean_contrast": "NR",
                    "standardized_contrast_state": "missing_or_empty_group",
                    "paired_contrast_50_cell_sensitivity": False,
                    "preexposed_code": code in PREEXPOSED_CODES,
                })
                continue
            if not np.all(np.isfinite(norm)) or not np.all(np.isfinite(malign)):
                raise ValueError("Non-finite expression value in paired comparator")
            normal_mean = float(np.mean(norm))
            malignant_mean = float(np.mean(malign))
            difference = malignant_mean - normal_mean
            if norm.size < 2:
                sd: object = "NR"
                zlike: object = "NR"
                state = "undefined_fewer_than_two_Normal_PC_cells"
            else:
                sd_value = float(np.std(norm, ddof=1))
                if not math.isfinite(sd_value):
                    sd = "NR"
                    zlike = "NR"
                    state = "undefined_nonfinite_Normal_PC_sd"
                elif sd_value == 0.0:
                    sd = 0.0
                    zlike = "NR"
                    state = "undefined_constant_Normal_PC_expression"
                else:
                    sd = sd_value
                    zlike = difference / sd_value
                    state = "descriptive_not_published_z_score"
            rows.append({
                "Hospital.Code": code,
                "marker": marker,
                "n_malignant_cells": int(malign.size),
                "n_within_code_Normal_PC_cells": int(norm.size),
                "paired_mean_difference_malignant_minus_Normal_PC": difference,
                "normal_reference_sd": sd,
                "standardized_mean_contrast": zlike,
                "standardized_contrast_state": state,
                "paired_contrast_50_cell_sensitivity": malign.size >= MIN_CELLS and norm.size >= MIN_CELLS,
                "preexposed_code": code in PREEXPOSED_CODES,
            })
    return rows


def evaluate_policy_set(
    analysis_set: str,
    healthy_codes: set[str],
    mm_codes: set[str],
    by_group_code: dict[tuple[str, str], list[dict[str, object]]],
) -> dict[str, object]:
    """Run code-held-out threshold comparisons for one prespecified code set."""
    fold_rows: list[dict[str, object]] = []
    code_rows: list[dict[str, object]] = []
    summary_rows: list[dict[str, object]] = []
    uncertain_rows: list[dict[str, object]] = []
    folds: dict[str, dict[str, int]] = {}
    if len(healthy_codes) < 4 or len(mm_codes) < 4:
        return {
            "available": False,
            "reason": "four-fold code-level comparison requires at least four nonempty codes in each group",
            "folds": folds,
            "cv_fold_rows": fold_rows,
            "cv_code_rows": code_rows,
            "cv_summary_rows": summary_rows,
            "uncertain_rows": uncertain_rows,
        }

    rng = np.random.default_rng(20261003)
    healthy_folds = make_folds(sorted(healthy_codes), 4, rng)
    mm_folds = make_folds(sorted(mm_codes), 4, rng)
    folds = {"healthy": healthy_folds, "mm": mm_folds}
    policies = ("pooled", "equal_code", "donor_envelope")
    budgets = (0.10, 0.05, 0.20, 0.50)

    for marker in MARKERS:
        score_cache = {
            (group, code): np.asarray([float(row[f"{marker}_per_10k"]) for row in rows], dtype=np.float64)
            for (group, code), rows in by_group_code.items()
        }
        for budget in budgets:
            for fold in range(4):
                train_codes = sorted(code for code in healthy_codes if healthy_folds[code] != fold)
                test_healthy = sorted(code for code in healthy_codes if healthy_folds[code] == fold)
                test_mm = sorted(code for code in mm_codes if mm_folds[code] == fold)
                training_primary = {code: score_cache[("healthy_normal_pc", code)] for code in train_codes}
                training_augmented = {
                    code: np.concatenate((
                        score_cache[("healthy_normal_pc", code)],
                        score_cache.get(("healthy_ambiguous_malignant_label", code), np.empty(0, dtype=np.float64)),
                    ))
                    for code in train_codes
                }
                training_max = max(float(np.max(values)) for values in training_primary.values())

                for policy in policies:
                    threshold, training_constraint = fit_thresholds(policy, training_primary, budget)
                    threshold_state = classify_threshold(threshold, training_max)
                    threshold_value = format_threshold(threshold)
                    test_groups = (
                        ("healthy_normal_pc", test_healthy, "heldout_healthy_normal_pc"),
                        ("mm_malignant", test_mm, "heldout_mm_malignant_label"),
                    )
                    for group, codes, label in test_groups:
                        positive_cells = 0
                        total_cells = 0
                        rates = []
                        for code in codes:
                            values = score_cache[(group, code)]
                            positives = threshold_positive(values, threshold)
                            rate = positives / values.size
                            rates.append(rate)
                            positive_cells += positives
                            total_cells += values.size
                            code_rows.append({
                                "analysis_set": analysis_set,
                                "marker": marker,
                                "budget": budget,
                                "policy": policy,
                                "fold": fold,
                                "holdout_group": label,
                                "Hospital.Code": code,
                                "n_cells": int(values.size),
                                "one_cell_fraction_resolution": 1.0 / values.size,
                                "positive_cells": positives,
                                "positive_fraction": rate,
                                "threshold_per_10k": threshold_value,
                                "threshold_state": threshold_state,
                                "preexposed_code": False,
                            })
                        fold_rows.append({
                            "analysis_set": analysis_set,
                            "marker": marker,
                            "budget": budget,
                            "policy": policy,
                            "fold": fold,
                            "threshold_per_10k": threshold_value,
                            "threshold_state": threshold_state,
                            "training_healthy_codes": len(train_codes),
                            "training_healthy_cells": sum(score_cache[("healthy_normal_pc", code)].size for code in train_codes),
                            "training_policy_constraint_fraction_at_or_above_threshold": training_constraint,
                            "heldout_group": label,
                            "heldout_codes": len(codes),
                            "heldout_cells": total_cells,
                            "heldout_positive_cells": positive_cells,
                            "heldout_cell_weighted_positive_fraction": positive_cells / total_cells if total_cells else "NR",
                            "heldout_equal_code_positive_fraction": float(np.mean(rates)) if rates else "NR",
                            "heldout_healthy_codes_above_normal_budget": sum(rate > budget for rate in rates) if label == "heldout_healthy_normal_pc" else "NR",
                            "heldout_mm_codes_with_coverage_above_illustrative_fraction": sum(rate > budget for rate in rates) if label == "heldout_mm_malignant_label" else "NR",
                        })

                    augmented_threshold, augmented_constraint = fit_thresholds(policy, training_augmented, budget)
                    augmented_max = max(float(np.max(values)) for values in training_augmented.values())
                    augmented_state = classify_threshold(augmented_threshold, augmented_max)
                    for group, codes, group_type in (
                        ("healthy_normal_pc", test_healthy, "heldout_healthy_normal_pc"),
                        ("mm_malignant", test_mm, "heldout_mm_malignant_label"),
                        ("healthy_ambiguous_malignant_label", test_healthy, "heldout_uncertain_healthy_malignant_label"),
                    ):
                        group_rates = []
                        positive_cells = 0
                        total_cells = 0
                        for code in codes:
                            values = score_cache.get((group, code), np.empty(0, dtype=np.float64))
                            if values.size == 0:
                                continue
                            positives = threshold_positive(values, augmented_threshold)
                            positive_cells += positives
                            total_cells += values.size
                            group_rates.append(positives / values.size)
                        uncertain_rows.append({
                            "analysis_set": analysis_set,
                            "marker": marker,
                            "budget": budget,
                            "policy": policy,
                            "fold": fold,
                            "sensitivity_group": group_type,
                            "n_codes": len(group_rates),
                            "n_cells": total_cells,
                            "one_cell_fraction_resolution_min": min((1.0 / score_cache[(group, code)].size for code in codes if score_cache.get((group, code), np.empty(0)).size), default="NR"),
                            "positive_cells": positive_cells,
                            "cell_weighted_positive_fraction": positive_cells / total_cells if total_cells else "NR",
                            "equal_code_positive_fraction": float(np.mean(group_rates)) if group_rates else "NR",
                            "augmented_training_threshold_per_10k": format_threshold(augmented_threshold),
                            "augmented_training_threshold_state": augmented_state,
                            "augmented_training_policy_constraint_fraction": augmented_constraint,
                        })

            for policy in policies:
                for group, label in (
                    ("heldout_healthy_normal_pc", "Healthy_Normal_PC"),
                    ("heldout_mm_malignant_label", "MM_Malignant_label"),
                ):
                    records = [
                        row for row in code_rows
                        if row["analysis_set"] == analysis_set and row["marker"] == marker
                        and row["budget"] == budget and row["policy"] == policy and row["holdout_group"] == group
                    ]
                    total_cells = sum(int(row["n_cells"]) for row in records)
                    positives = sum(int(row["positive_cells"]) for row in records)
                    rates = [float(row["positive_fraction"]) for row in records]
                    resolutions = [float(row["one_cell_fraction_resolution"]) for row in records]
                    summary_rows.append({
                        "analysis_set": analysis_set,
                        "marker": marker,
                        "budget": budget,
                        "policy": policy,
                        "heldout_group": label,
                        "heldout_codes": len(records),
                        "heldout_cells": total_cells,
                        "cell_weighted_positive_fraction": positives / total_cells if total_cells else "NR",
                        "equal_code_positive_fraction": float(np.mean(rates)) if rates else "NR",
                        "minimum_code_positive_fraction": min(rates) if rates else "NR",
                        "median_code_positive_fraction": float(np.median(rates)) if rates else "NR",
                        "maximum_code_positive_fraction": max(rates) if rates else "NR",
                        "minimum_one_cell_fraction_resolution": min(resolutions) if resolutions else "NR",
                        "maximum_one_cell_fraction_resolution": max(resolutions) if resolutions else "NR",
                        "healthy_codes_above_normal_budget": sum(rate > budget for rate in rates) if group == "heldout_healthy_normal_pc" else "NR",
                        "mm_codes_with_coverage_above_illustrative_fraction": sum(rate > budget for rate in rates) if group == "heldout_mm_malignant_label" else "NR",
                        "precision_note": "observed code-level outcomes and fold variation; no population interval or independent replication",
                    })

    return {
        "available": True,
        "reason": "four-fold code-level comparison completed",
        "folds": folds,
        "cv_fold_rows": fold_rows,
        "cv_code_rows": code_rows,
        "cv_summary_rows": summary_rows,
        "uncertain_rows": uncertain_rows,
    }


def main() -> None:
    checks = {
        str(ANNOTATION): (sha256(ANNOTATION), EXPECTED_WORKBOOK_SHA),
        str(PILOT_MANIFEST): (sha256(PILOT_MANIFEST), EXPECTED_PILOT_MANIFEST_SHA),
        str(DOWNLOAD_MANIFEST): (sha256(DOWNLOAD_MANIFEST), EXPECTED_DOWNLOAD_MANIFEST_SHA),
        str(PROTOCOL): (sha256(PROTOCOL), EXPECTED_PROTOCOL_SHA),
    }
    protocol_sidecar = PROTOCOL_HASH.read_text(encoding="utf-8").split()[0]
    if protocol_sidecar != EXPECTED_PROTOCOL_SHA:
        raise ValueError("Selected protocol SHA-256 sidecar mismatch")
    failed = {path: values for path, values in checks.items() if values[0] != values[1]}
    if failed:
        raise ValueError(f"Frozen input hash check failed: {failed}")

    with PILOT_SUMMARY.open("r", encoding="utf-8") as handle:
        selection_summary = json.load(handle)
    healthy_codes = set(selection_summary["selected_healthy_codes"])
    all_mm_codes = set(selection_summary["selected_mm_codes"])
    primary_mm_codes = all_mm_codes - PREEXPOSED_CODES
    selected_codes = healthy_codes | all_mm_codes
    with DOWNLOAD_MANIFEST.open("r", encoding="utf-8", newline="") as handle:
        file_rows = list(csv.DictReader(handle, delimiter="\t"))
    if len(file_rows) != 86 or len({r["GSM"] for r in file_rows}) != 86:
        raise ValueError(f"Expected 86 unique downloaded GSM files, found {len(file_rows)}")
    file_by_batch: dict[str, dict[str, str]] = {}
    for row in file_rows:
        path = DATA / row["filename"]
        if not path.is_file() or path.stat().st_size != int(row["bytes"]) or sha256(path) != row["sha256"]:
            raise ValueError(f"Downloaded count-table hash/size mismatch: {path.name}")
        if row["Amp.Batch.ID"] in file_by_batch:
            raise ValueError(f"Duplicate MARS GEO count file for batch {row['Amp.Batch.ID']}")
        file_by_batch[row["Amp.Batch.ID"]] = row
    batch_ids = set(file_by_batch)
    annotations, annotated_counts = read_selected_annotations(selected_codes, batch_ids)
    annotations_by_batch: dict[str, set[str]] = defaultdict(set)
    for cell_id, annotation in annotations.items():
        annotations_by_batch[annotation["Amp.Batch.ID"]].add(cell_id)

    cell_rows: list[dict[str, object]] = []
    file_audit_rows: list[dict[str, object]] = []
    seen_cell_ids: set[str] = set()
    for file_index, batch_id in enumerate(sorted(file_by_batch), start=1):
        record = file_by_batch[batch_id]
        path = DATA / record["filename"]
        ids, positions, totals, marker_values, gene_count = parse_count_file(path, set(annotations))
        selected_positions = [p for p, cell_id in enumerate(ids) if cell_id in annotations]
        keyed_matches = exact_batch_matches(ids, annotations, batch_id)
        matched_same_batch = [position for position, _, _ in keyed_matches]
        batch_mismatch = [p for p in selected_positions if annotations[ids[p]]["Amp.Batch.ID"] != batch_id]
        for pos in matched_same_batch:
            cell_id = ids[pos]
            if cell_id in seen_cell_ids:
                raise ValueError(f"Selected cell appears in multiple matrices: {cell_id}")
            seen_cell_ids.add(cell_id)
            annotation = annotations[cell_id]
            total = int(totals[pos])
            if total <= 0:
                raise ValueError(f"Selected cell has no captured molecules: {cell_id}")
            cell = {
                "cell_id": cell_id,
                "GSM": record["GSM"],
                "GEO_series": record["GEO_series"],
                **annotation,
                "total_mRNA_molecule_counts": total,
                "preexposed_code": annotation["Hospital.Code"] in PREEXPOSED_CODES,
            }
            for marker in MARKERS:
                raw = int(marker_values[marker][pos])
                if raw > total:
                    raise ValueError(f"{marker} molecules exceed all-gene total for {cell_id}")
                cell[f"{marker}_molecules"] = raw
                cell[f"{marker}_per_10k"] = (raw / total) * 10000.0
            cell_rows.append(cell)
        file_audit_rows.append({
            "GSM": record["GSM"],
            "GEO_series": record["GEO_series"],
            "Amp.Batch.ID": batch_id,
            "file": record["filename"],
            "matrix_cell_columns": len(ids),
            "gene_rows": gene_count,
            "selected_annotation_cell_ids_for_batch": len(annotations_by_batch[batch_id]),
            "exact_joined_selected_same_batch_cells": len(matched_same_batch),
            "cross_batch_annotation_matches": len(batch_mismatch),
            "count_columns_without_selected_annotation": len(ids) - len(selected_positions),
            "table_sha256": record["sha256"],
            "preexposed_GSM": record["GSM"] == "GSM9189168",
        })
        if file_index % 10 == 0 or file_index == len(file_by_batch):
            print(f"parsed {file_index}/{len(file_by_batch)} count tables; joined selected cells={len(seen_cell_ids)}", flush=True)

    cells_fields = (
        "cell_id", "GSM", "GEO_series", "Amp.Batch.ID", "Hospital.Code", "Biopsy.Sequence",
        "Disease", "Tissue", "Method", "Populations", "Seq.Batch.ID", "analysis_group",
        "total_mRNA_molecule_counts", "TNFRSF17_molecules", "TNFRSF17_per_10k",
        "GPRC5D_molecules", "GPRC5D_per_10k", "SLAMF7_molecules", "SLAMF7_per_10k",
        "preexposed_code",
    )
    write_tsv(RESULTS / "external_pilot_selected_cells.tsv", cells_fields, cell_rows)

    code_rows, by_group_code = metric_rows(cell_rows, selected_codes)
    flow_rows = []
    for code in sorted(selected_codes):
        for group in (
            "healthy_normal_pc", "healthy_ambiguous_malignant_label", "mm_malignant", "mm_normal_pc_prior_art_reference"
        ):
            annotation_n = annotated_counts.get((group, code), 0)
            matched_n = len(by_group_code.get((group, code), []))
            flow_rows.append({
                "Hospital.Code": code,
                "analysis_group": group,
                "workbook_eligible_cells": annotation_n,
                "exact_joined_cells": matched_n,
                "meets_50_cell_sensitivity": matched_n >= MIN_CELLS,
                "preexposed_code": code in PREEXPOSED_CODES,
                "exclusion_or_limit": "prior marker exposure; excluded from primary metrics" if code in PREEXPOSED_CODES else (
                    "below 50-cell sensitivity screen; retained in primary" if 0 < matched_n < MIN_CELLS else ("no exact-joined cells in this group" if matched_n == 0 else "none")
                ),
            })

    write_tsv(
        RESULTS / "external_pilot_file_join_audit.tsv",
        ("GSM", "GEO_series", "Amp.Batch.ID", "file", "matrix_cell_columns", "gene_rows",
         "selected_annotation_cell_ids_for_batch", "exact_joined_selected_same_batch_cells",
         "cross_batch_annotation_matches", "count_columns_without_selected_annotation", "table_sha256", "preexposed_GSM"),
        file_audit_rows,
    )
    write_tsv(
        RESULTS / "external_pilot_code_flow.tsv",
        ("Hospital.Code", "analysis_group", "workbook_eligible_cells", "exact_joined_cells", "meets_50_cell_sensitivity", "preexposed_code", "exclusion_or_limit"),
        flow_rows,
    )
    write_tsv(
        RESULTS / "external_pilot_code_metrics.tsv",
        ("analysis_group", "Hospital.Code", "marker", "n_exact_joined_cells", "meets_50_cell_sensitivity", "preexposed_code",
         "median_total_mRNA_molecules", "mean_raw_molecules", "raw_detection_fraction", "mean_per_10k", "median_per_10k", "q90_per_10k"),
        code_rows,
    )

    primary_healthy_codes = {
        code for code in healthy_codes
        if len(by_group_code.get(("healthy_normal_pc", code), [])) > 0
    }
    primary_mm_codes_nonempty = {
        code for code in primary_mm_codes
        if len(by_group_code.get(("mm_malignant", code), [])) > 0
    }
    sensitivity_healthy_codes = {
        code for code in primary_healthy_codes
        if len(by_group_code.get(("healthy_normal_pc", code), [])) >= MIN_CELLS
    }
    sensitivity_mm_codes = {
        code for code in primary_mm_codes_nonempty
        if len(by_group_code.get(("mm_malignant", code), [])) >= MIN_CELLS
    }
    primary_evaluation = evaluate_policy_set(
        "all_nonempty_primary", primary_healthy_codes, primary_mm_codes_nonempty, by_group_code
    )
    sensitivity_evaluation = evaluate_policy_set(
        "min50_cells_sensitivity", sensitivity_healthy_codes, sensitivity_mm_codes, by_group_code
    )
    evaluations = (primary_evaluation, sensitivity_evaluation)
    cv_fold_rows = [row for result in evaluations for row in result["cv_fold_rows"]]
    cv_code_rows = [row for result in evaluations for row in result["cv_code_rows"]]
    cv_summary_rows = [row for result in evaluations for row in result["cv_summary_rows"]]
    uncertain_rows = [row for result in evaluations for row in result["uncertain_rows"]]
    folds = {
        "all_nonempty_primary": primary_evaluation["folds"],
        "min50_cells_sensitivity": sensitivity_evaluation["folds"],
    }

    paired_rows = paired_contrasts(by_group_code, all_mm_codes)
    write_tsv(
        RESULTS / "external_pilot_paired_contrasts.tsv",
        ("Hospital.Code", "marker", "n_malignant_cells", "n_within_code_Normal_PC_cells",
         "paired_mean_difference_malignant_minus_Normal_PC", "normal_reference_sd", "standardized_mean_contrast",
         "standardized_contrast_state", "paired_contrast_50_cell_sensitivity", "preexposed_code"),
        paired_rows,
    )
    write_tsv(
        RESULTS / "external_pilot_cv_thresholds.tsv",
        ("analysis_set", "marker", "budget", "policy", "fold", "threshold_per_10k", "training_healthy_codes", "training_healthy_cells",
         "threshold_state", "training_policy_constraint_fraction_at_or_above_threshold", "heldout_group", "heldout_codes", "heldout_cells",
         "heldout_positive_cells", "heldout_cell_weighted_positive_fraction", "heldout_equal_code_positive_fraction",
         "heldout_healthy_codes_above_normal_budget", "heldout_mm_codes_with_coverage_above_illustrative_fraction"),
        cv_fold_rows,
    )
    write_tsv(
        RESULTS / "external_pilot_cv_by_code.tsv",
        ("analysis_set", "marker", "budget", "policy", "fold", "holdout_group", "Hospital.Code", "n_cells",
         "one_cell_fraction_resolution", "positive_cells", "positive_fraction", "threshold_per_10k", "threshold_state", "preexposed_code"),
        cv_code_rows,
    )
    write_tsv(
        RESULTS / "external_pilot_cv_summary.tsv",
        ("analysis_set", "marker", "budget", "policy", "heldout_group", "heldout_codes", "heldout_cells",
         "cell_weighted_positive_fraction", "equal_code_positive_fraction", "minimum_code_positive_fraction",
         "median_code_positive_fraction", "maximum_code_positive_fraction", "minimum_one_cell_fraction_resolution",
         "maximum_one_cell_fraction_resolution", "healthy_codes_above_normal_budget",
         "mm_codes_with_coverage_above_illustrative_fraction", "precision_note"),
        cv_summary_rows,
    )
    write_tsv(
        RESULTS / "external_pilot_uncertain_label_sensitivity.tsv",
        ("analysis_set", "marker", "budget", "policy", "fold", "sensitivity_group", "n_codes", "n_cells",
         "one_cell_fraction_resolution_min", "positive_cells", "cell_weighted_positive_fraction", "equal_code_positive_fraction",
         "augmented_training_threshold_per_10k", "augmented_training_threshold_state", "augmented_training_policy_constraint_fraction"),
        uncertain_rows,
    )

    primary_joined_cells = sum(
        len(by_group_code.get(("healthy_normal_pc", code), [])) for code in primary_healthy_codes
    ) + sum(len(by_group_code.get(("mm_malignant", code), [])) for code in primary_mm_codes_nonempty)
    preexposed_joined_cells = sum(
        1 for row in cell_rows if row["Hospital.Code"] in PREEXPOSED_CODES
    )
    summary = {
        "protocol_sha256": EXPECTED_PROTOCOL_SHA,
        "prior_protocol_v1_sha256": "799a246abb8fe939e4937fd5bd9f97eaf04e6a6222b8aee56d720b3544c8478c",
        "earlier_reported_lock_sha256": "5f5c961d…f793cd2 (not yet located as a preserved file)",
        "workbook_sha256": EXPECTED_WORKBOOK_SHA,
        "locked_pilot_manifest_sha256": EXPECTED_PILOT_MANIFEST_SHA,
        "download_manifest_sha256": EXPECTED_DOWNLOAD_MANIFEST_SHA,
        "selected_GSM_tables": len(file_rows),
        "selected_batch_ids": len(file_by_batch),
        "selected_annotation_cell_ids": len(annotations),
        "exact_joined_unique_cells": len(cell_rows),
        "exact_joined_primary_healthy_and_MM_cells_excluding_preexposed_code": primary_joined_cells,
        "exact_joined_cells_in_preexposed_code_excluded_from_primary": preexposed_joined_cells,
        "primary_healthy_codes_selected": len(healthy_codes),
        "primary_healthy_codes_nonempty_exact_join": len(primary_healthy_codes),
        "primary_mm_codes_selected": len(primary_mm_codes),
        "primary_mm_codes_nonempty_exact_join_excluding_preexposed": len(primary_mm_codes_nonempty),
        "primary_healthy_codes_meeting_50_cell_sensitivity": len(sensitivity_healthy_codes),
        "primary_mm_codes_meeting_50_cell_sensitivity": len(sensitivity_mm_codes),
        "primary_four_fold_comparison_available": primary_evaluation["available"],
        "primary_four_fold_comparison_note": primary_evaluation["reason"],
        "50_cell_sensitivity_four_fold_comparison_available": sensitivity_evaluation["available"],
        "50_cell_sensitivity_four_fold_comparison_note": sensitivity_evaluation["reason"],
        "preexposed_codes_excluded_from_primary_metrics": sorted(PREEXPOSED_CODES & selected_codes),
        "counts_by_source_group": dict(sorted(Counter(row["analysis_group"] for row in cell_rows).items())),
        "matrix_total_definition": "sum of all nonnegative integer gene-row molecule counts per cell",
        "normalization": "10,000 * marker mRNA molecule count / all-gene-row molecule sum per cell",
        "cv_fold_assignments": folds,
        "allowed_interpretation": "small, within-GSE305971 code-held-out threshold-policy pilot only; no clinical or population-level claim",
        "data_gate_note": "Four folds require at least four nonempty putative codes per group; this is an analysis-availability condition, not evidence of precision. Every nonempty code remains in primary descriptive results; the 50-cell screen is sensitivity only.",
        "outputs": [
            "external_pilot_selected_cells.tsv", "external_pilot_file_join_audit.tsv", "external_pilot_code_flow.tsv",
            "external_pilot_code_metrics.tsv", "external_pilot_paired_contrasts.tsv",
            "external_pilot_cv_thresholds.tsv", "external_pilot_cv_by_code.tsv", "external_pilot_cv_summary.tsv",
            "external_pilot_uncertain_label_sensitivity.tsv",
        ],
    }
    (RESULTS / "external_pilot_result_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8"
    )
    print(json.dumps({k: summary[k] for k in (
        "selected_GSM_tables", "selected_annotation_cell_ids", "exact_joined_unique_cells",
        "primary_healthy_codes_nonempty_exact_join", "primary_mm_codes_nonempty_exact_join_excluding_preexposed",
        "primary_four_fold_comparison_available", "50_cell_sensitivity_four_fold_comparison_available", "counts_by_source_group",
    )}, sort_keys=True))


if __name__ == "__main__":
    main()
