"""Create auditable sample and source-identifier inventories for GSE305971."""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
NEXT = HERE.parent
SOURCE = NEXT / "data" / "GSE305971_cells_annotation_for_geo_ver_20250306.xlsx"
OUT_DIR = HERE / "results"
BY_BATCH = OUT_DIR / "external_annotation_batch_inventory.tsv"
BY_HOSPITAL_BIOPSY = OUT_DIR / "external_annotation_by_hospital_biopsy.tsv"
SUMMARY = OUT_DIR / "external_annotation_identity_audit.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalized(value: object) -> str:
    if value is None or not str(value).strip():
        return "NR"
    return str(value).strip()


def compact_json(values: Counter) -> str:
    return json.dumps(dict(sorted(values.items())), ensure_ascii=False, sort_keys=True)


def main() -> None:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    columns = {name: index for index, name in enumerate(header)}
    required = (
        "index", "Amp.Batch.ID", "Hospital.Code", "Biopsy.Sequence", "Populations",
        "Disease", "Tissue", "Method", "Seq.Batch.ID",
    )
    missing = [name for name in required if name not in columns]
    if missing:
        raise ValueError(f"Annotation columns missing: {missing}")

    batch_records: dict[str, dict[str, object]] = {}
    pair_records: dict[tuple[str, str], dict[str, object]] = {}
    disease_population_cells: Counter[tuple[str, str, str]] = Counter()
    hospital_to_sequences: dict[str, set[str]] = defaultdict(set)
    source_ids: set[str] = set()
    total_cells = 0

    for row_number, row in enumerate(rows, start=2):
        if len(row) != len(header):
            raise ValueError(f"Malformed annotation row {row_number}")
        cell_id = normalized(row[columns["index"]])
        if cell_id == "NR" or cell_id in source_ids:
            raise ValueError(f"Missing or duplicate cell ID at row {row_number}: {cell_id}")
        source_ids.add(cell_id)
        values = {name: normalized(row[columns[name]]) for name in required[1:]}
        if values["Amp.Batch.ID"] == "NR":
            raise ValueError(f"Missing Amp.Batch.ID at row {row_number}")

        batch_id = values["Amp.Batch.ID"]
        batch = batch_records.setdefault(
            batch_id,
            {
                "cells": 0,
                "hospital": Counter(),
                "biopsy": Counter(),
                "disease": Counter(),
                "tissue": Counter(),
                "method": Counter(),
                "seq_batch": Counter(),
                "population": Counter(),
                "hospital_biopsy_disease": Counter(),
                "disease_tissue_method_population": Counter(),
            },
        )
        batch["cells"] += 1
        for field, key in (
            ("Hospital.Code", "hospital"), ("Biopsy.Sequence", "biopsy"),
            ("Disease", "disease"), ("Tissue", "tissue"), ("Method", "method"),
            ("Seq.Batch.ID", "seq_batch"), ("Populations", "population"),
        ):
            batch[key][values[field]] += 1
        batch["hospital_biopsy_disease"][
            "|".join((values["Hospital.Code"], values["Biopsy.Sequence"], values["Disease"]))
        ] += 1
        batch["disease_tissue_method_population"][
            "|".join((values["Disease"], values["Tissue"], values["Method"], values["Populations"]))
        ] += 1

        hospital = values["Hospital.Code"]
        biopsy = values["Biopsy.Sequence"]
        hospital_to_sequences[hospital].add(biopsy)
        pair_key = (hospital, biopsy)
        pair = pair_records.setdefault(
            pair_key,
            {
                "cells": 0,
                "disease": Counter(),
                "tissue": Counter(),
                "method": Counter(),
                "population": Counter(),
                "amp_batch": Counter(),
                "seq_batch": Counter(),
                "disease_tissue_method_population": Counter(),
            },
        )
        pair["cells"] += 1
        pair["disease"][values["Disease"]] += 1
        pair["tissue"][values["Tissue"]] += 1
        pair["method"][values["Method"]] += 1
        pair["population"][values["Populations"]] += 1
        pair["amp_batch"][batch_id] += 1
        pair["seq_batch"][values["Seq.Batch.ID"]] += 1
        pair["disease_tissue_method_population"][
            "|".join((values["Disease"], values["Tissue"], values["Method"], values["Populations"]))
        ] += 1
        disease_population_cells[(values["Disease"], values["Tissue"], values["Populations"])] += 1
        total_cells += 1

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    batch_fields = (
        "Amp.Batch.ID", "annotated_cells", "Hospital.Code_counts", "Biopsy.Sequence_counts",
        "Disease_counts", "Tissue_counts", "Method_counts", "Seq.Batch.ID_counts",
        "Populations_counts", "Hospital_Biopsy_Disease_counts", "Disease_Tissue_Method_Population_counts",
    )
    with BY_BATCH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=batch_fields, delimiter="\t", quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        for batch_id, record in sorted(batch_records.items()):
            writer.writerow({
                "Amp.Batch.ID": batch_id,
                "annotated_cells": record["cells"],
                "Hospital.Code_counts": compact_json(record["hospital"]),
                "Biopsy.Sequence_counts": compact_json(record["biopsy"]),
                "Disease_counts": compact_json(record["disease"]),
                "Tissue_counts": compact_json(record["tissue"]),
                "Method_counts": compact_json(record["method"]),
                "Seq.Batch.ID_counts": compact_json(record["seq_batch"]),
                "Populations_counts": compact_json(record["population"]),
                "Hospital_Biopsy_Disease_counts": compact_json(record["hospital_biopsy_disease"]),
                "Disease_Tissue_Method_Population_counts": compact_json(record["disease_tissue_method_population"]),
            })

    pair_fields = (
        "Hospital.Code", "Biopsy.Sequence", "annotated_cells", "Disease_counts", "Tissue_counts",
        "Method_counts", "Populations_counts", "Amp.Batch.ID_counts", "Seq.Batch.ID_counts",
        "Disease_Tissue_Method_Population_counts", "identity_interpretation",
    )
    with BY_HOSPITAL_BIOPSY.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=pair_fields, delimiter="\t", quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        for (hospital, biopsy), record in sorted(pair_records.items()):
            writer.writerow({
                "Hospital.Code": hospital,
                "Biopsy.Sequence": biopsy,
                "annotated_cells": record["cells"],
                "Disease_counts": compact_json(record["disease"]),
                "Tissue_counts": compact_json(record["tissue"]),
                "Method_counts": compact_json(record["method"]),
                "Populations_counts": compact_json(record["population"]),
                "Amp.Batch.ID_counts": compact_json(record["amp_batch"]),
                "Seq.Batch.ID_counts": compact_json(record["seq_batch"]),
                "Disease_Tissue_Method_Population_counts": compact_json(record["disease_tissue_method_population"]),
                "identity_interpretation": "source sample key; use Hospital.Code conservatively for repeated-biopsy grouping, but participant semantics are not validated",
            })

    disease_discordant_pairs = sum(len(record["disease"]) > 1 for record in pair_records.values())
    tissue_mixed_pairs = sum(len(record["tissue"]) > 1 for record in pair_records.values())
    method_mixed_pairs = sum(len(record["method"]) > 1 for record in pair_records.values())
    summary = {
        "annotation_file": str(SOURCE),
        "annotation_file_bytes": SOURCE.stat().st_size,
        "annotation_file_sha256": sha256(SOURCE),
        "annotation_rows": total_cells,
        "unique_cell_ids": len(source_ids),
        "amp_batch_ids": len(batch_records),
        "hospital_code_values": len(hospital_to_sequences),
        "hospital_biopsy_sequence_pairs": len(pair_records),
        "hospital_codes_with_multiple_biopsy_sequences": sum(len(values) > 1 for values in hospital_to_sequences.values()),
        "max_sequences_per_hospital_code": max((len(values) for values in hospital_to_sequences.values()), default=0),
        "hospital_biopsy_pairs_with_multiple_disease_values": disease_discordant_pairs,
        "hospital_biopsy_pairs_with_multiple_tissue_values": tissue_mixed_pairs,
        "hospital_biopsy_pairs_with_multiple_method_values": method_mixed_pairs,
        "cells_by_disease_tissue_population": {
            "|".join(key): count for key, count in sorted(disease_population_cells.items())
        },
        "identity_warning": "The annotation source provides Hospital.Code and Biopsy.Sequence fields but does not define them as a validated person identifier. Pair-level rows are source keys only; do not claim unique patient counts or independence without a primary-source crosswalk.",
        "superseded_output": "external_annotation_by_biopsy.tsv groups on Biopsy.Sequence alone and merges distinct Hospital.Code values.",
        "batch_output": str(BY_BATCH),
        "hospital_biopsy_output": str(BY_HOSPITAL_BIOPSY),
    }
    SUMMARY.write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({
        "annotation_rows": total_cells,
        "amp_batch_ids": len(batch_records),
        "hospital_code_values": len(hospital_to_sequences),
        "hospital_biopsy_pairs": len(pair_records),
        "hospital_codes_with_multiple_sequences": summary["hospital_codes_with_multiple_biopsy_sequences"],
        "disease_discordant_hospital_biopsy_pairs": disease_discordant_pairs,
        "mixed_tissue_hospital_biopsy_pairs": tissue_mixed_pairs,
        "mixed_method_hospital_biopsy_pairs": method_mixed_pairs,
        "disease_tissue_population_cells": summary["cells_by_disease_tissue_population"],
        "outputs": [str(BY_BATCH), str(BY_HOSPITAL_BIOPSY), str(SUMMARY)],
    }, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
