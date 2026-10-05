import pathlib,json,collections,csv
import numpy as np
import scipy.io as sio
import openpyxl
root=pathlib.Path(__file__).resolve().parent
report={'spreadsheets':{},'matlab':{}}
for p in root.glob('*.xlsx'):
    wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
    summaries=[]
    for s in wb:
        rows=list(s.iter_rows(values_only=True))
        nonempty=[list(r[:12]) for r in rows if any(v is not None for v in r)]
        summaries.append({'sheet':s.title,'rows':s.max_row,'columns':s.max_column,'first_nonempty_rows':nonempty[:8]})
    report['spreadsheets'][p.name]=summaries;wb.close()
for p in (root/'trail_source').rglob('*.mat'):
    data=sio.loadmat(p)
    summary={}
    for k,v in data.items():
        if k.startswith('__'):continue
        obj={'shape':list(v.shape),'dtype':str(v.dtype)}
        if np.issubdtype(v.dtype,np.number):
            obj.update({'finite':int(np.isfinite(v).sum()),'nan':int(np.isnan(v).sum()),'first_values':v.ravel()[:8].tolist()})
        summary[k]=obj
    report['matlab'][str(p.relative_to(root))]=summary
(root/'numeric_inventory.json').write_text(json.dumps(report,indent=2,default=str),encoding='utf-8')
for name,sheets in report['spreadsheets'].items():
    print('WORKBOOK',name)
    for s in sheets:print(json.dumps(s,default=str))
for name,data in report['matlab'].items():
    if '/Figure 2/' in name.replace('\\','/') and 'FRET' in name: print('TRAIL',name,json.dumps(data))
